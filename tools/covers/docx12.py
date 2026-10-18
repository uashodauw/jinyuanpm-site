import docx,re,os,io,json
from docx.oxml.ns import qn
from PIL import Image
SITE='/workspace/jinyuanpm-site'; slug='personal-agent-from-tool-to-collaborator'
d=docx.Document('/workspace/article12/source.docx'); D=f'{SITE}/assets/articles/{slug}'; os.makedirs(D,exist_ok=True)
def runs_md(p):
    out=''
    for r in p.runs:
        t=r.text
        if not t: continue
        if r.bold and t.strip():
            lead=t[:len(t)-len(t.lstrip())]; tr=t[len(t.rstrip()):]; t=lead+'**'+t.strip()+'**'+tr
        out+=t
    return re.sub(r'\*\*\*\*','',out).strip()
def cell_md(c): return ' '.join(runs_md(p) for p in c.paragraphs if p.text.strip()).replace('|','／')
items=[]
for el in d.element.body.iterchildren():
    tag=el.tag.split('}')[1]
    if tag=='p':
        p=docx.text.paragraph.Paragraph(el,d)
        for b in el.findall('.//'+qn('a:blip')): items.append(('img',d.part.related_parts[b.get(qn('r:embed'))]))
        if p.text.strip(): items.append((p.style.name,p))
    elif tag=='tbl': items.append(('tbl',docx.table.Table(el,d)))
out=[]; imgno=0; i=0; skip=False; removed=[]
while i<len(items):
    k,v=items[i]; i+=1
    if k=='Heading 1':
        h=v.text.strip(); skip=bool(re.match(r'^(七|八)\s',h))
        if skip: removed.append(h); continue
    if skip: continue
    if k=='img':
        imgno+=1
        im=Image.open(io.BytesIO(v.blob)); im=im.convert('RGB')
        cap=''
        if i<len(items) and items[i][0]=='Caption': cap=re.sub(r'\s+',' ',items[i][1].text.strip()); i+=1
        x=im.resize((1200,round(im.height*1200/im.width)),Image.LANCZOS) if im.width>1200 else im
        fn=f'{imgno:02d}.webp'; x.save(f'{D}/{fn}','WEBP',quality=78,method=6)
        print('img',im.size,'->',x.size,os.path.getsize(f'{D}/{fn}')//1024,'KB',cap)
        out.append(f'![{cap}](assets/articles/{slug}/{fn} "{cap}"){{{x.width}x{x.height}}}'); continue
    if k=='tbl':
        rows=[[cell_md(c) for c in r.cells] for r in v.rows]
        out.append('\n'.join(['| '+' | '.join(rows[0])+' |','|'+'---|'*len(rows[0])]+['| '+' | '.join(r)+' |' for r in rows[1:]])); continue
    t=runs_md(v)
    if k=='Title': title=v.text.strip(); continue
    if k=='Heading 1': out.append('## '+v.text.strip()); continue
    if k=='Heading 2': out.append('### '+v.text.strip()); continue
    if k=='Caption': out.append('*'+t+'*'); continue
    if k.startswith('List'): out.append('- '+t); continue
    out.append(t)
md='\n\n'.join(out)
open(f'site/{slug}.md','w').write(md+'\n')
words=len(re.findall('[\u4e00-\u9fff]',md))+len(re.findall(r'[A-Za-z0-9]+',md))
print(title,'| removed',removed,'| words',words,round(words/400),'| placeholders',re.findall(r'需补充|待补充|TODO|水印|占位',md))
